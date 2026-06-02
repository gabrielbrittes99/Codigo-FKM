"""
Módulo de Load (Carregamento) de Dados
Responsável por salvar DataFrames em arquivos Excel com a formatação padrão da empresa.
"""

import logging

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

logger = logging.getLogger(__name__)


def preparar_resumos(df_filial: pd.DataFrame):
    """
    Separa e agrupa os dados aplicando a regra de negócio para ignorar Arla e hodômetros inválidos.
    """
    mask_arla = df_filial["Combustivel"].str.contains("Arla", case=False, na=False)
    df_sem_arla = df_filial[~mask_arla].copy()
    df_arla = df_filial[mask_arla].copy()

    resumo_combustivel = None
    if not df_sem_arla.empty:
        df_km = df_sem_arla[
            (df_sem_arla["Hodometro/Horimetro"] > 0)
            & (df_sem_arla["Hodometro/Horimetro anterior"] > 0)
        ].copy()

        resumo_custos = (
            df_sem_arla.groupby("Placa")
            .agg({"Litragem": "sum", "Valor total": "sum"})
            .reset_index()
        )
        resumo_kms = (
            df_km.groupby("Placa")
            .agg(
                {
                    "Hodometro/Horimetro anterior": "min",
                    "Hodometro/Horimetro": "max",
                }
            )
            .reset_index()
        )
        resumo_combustivel = pd.merge(resumo_custos, resumo_kms, on="Placa", how="left")

        resumo_combustivel.columns = [
            "Rótulos de Linha",
            "Soma de Litragem",
            "Soma de Valor total",
            "Min. de Hodometro/Horimetro anterior",
            "Máx. de Hodometro/Horimetro",
        ]

        total_combustivel = pd.DataFrame(
            [
                {
                    "Rótulos de Linha": "Total Geral",
                    "Soma de Litragem": resumo_combustivel["Soma de Litragem"].sum(),
                    "Soma de Valor total": resumo_combustivel[
                        "Soma de Valor total"
                    ].sum(),
                    "Min. de Hodometro/Horimetro anterior": resumo_combustivel[
                        "Min. de Hodometro/Horimetro anterior"
                    ].min(),
                    "Máx. de Hodometro/Horimetro": resumo_combustivel[
                        "Máx. de Hodometro/Horimetro"
                    ].max(),
                }
            ]
        )
        resumo_combustivel = pd.concat(
            [resumo_combustivel, total_combustivel], ignore_index=True
        )

    resumo_arla = None
    if not df_arla.empty:
        resumo_arla = df_arla.groupby("Placa").agg({"Valor total": "sum"}).reset_index()
        resumo_arla.columns = ["Rótulos de Linha", "Soma de Valor total"]
        total_arla = pd.DataFrame(
            [
                {
                    "Rótulos de Linha": "Total Geral",
                    "Soma de Valor total": resumo_arla["Soma de Valor total"].sum(),
                }
            ]
        )
        resumo_arla = pd.concat([resumo_arla, total_arla], ignore_index=True)

    resumo_posto = (
        df_sem_arla.groupby("Estabelecimento")
        .agg({"Litragem": "sum", "Valor total": "sum"})
        .reset_index()
    )
    resumo_posto.columns = [
        "Rótulos de Linha",
        "Soma de Litragem",
        "Soma de Valor total",
    ]
    total_posto = pd.DataFrame(
        [
            {
                "Rótulos de Linha": "Total Geral",
                "Soma de Litragem": resumo_posto["Soma de Litragem"].sum(),
                "Soma de Valor total": resumo_posto["Soma de Valor total"].sum(),
            }
        ]
    )
    resumo_posto = pd.concat([resumo_posto, total_posto], ignore_index=True)

    return resumo_combustivel, resumo_arla, resumo_posto


def salvar_resumos_filial_excel(
    df_filial: pd.DataFrame, nome_filial: str, caminho_saida: str
) -> bool:
    """
    Gera o arquivo Excel com as abas separadas por tipo de combustível
    e aplica a formatação visual padrão.

    Args:
        df_filial (pd.DataFrame): Dados da filial já transformados.
        nome_filial (str): Nome da filial.
        caminho_saida (str): Caminho completo onde o Excel será salvo.

    Returns:
        bool: True se salvo com sucesso, False caso contrário.
    """
    logger.info(f"Iniciando a geração do Excel para a filial: {nome_filial}")

    try:
        resumo_combustivel, resumo_arla, resumo_posto = preparar_resumos(df_filial)

        # ==================== SALVAR EXCEL ====================
        logger.debug(f"Salvando planilhas no arquivo: {caminho_saida}")
        with pd.ExcelWriter(caminho_saida, engine="openpyxl") as writer:
            df_filial.to_excel(writer, sheet_name="Dados Brutos", index=False)
            if resumo_combustivel is not None:
                resumo_combustivel.to_excel(
                    writer, sheet_name="Combustível (Vários itens)", index=False
                )
            if resumo_arla is not None:
                resumo_arla.to_excel(writer, sheet_name="Arla 32", index=False)
            resumo_posto.to_excel(writer, sheet_name="Resumo por Posto", index=False)

        # ==================== APLICAR FORMATAÇÃO ====================
        logger.debug("Aplicando formatações visuais (Cores, Fontes e Numéricos)")
        wb = load_workbook(caminho_saida)
        cor_amarelo = PatternFill(
            start_color="FFFF00", end_color="FFFF00", fill_type="solid"
        )
        cor_cinza = PatternFill(
            start_color="D3D3D3", end_color="D3D3D3", fill_type="solid"
        )
        fonte_negrito = Font(bold=True)
        alinhamento_centro = Alignment(horizontal="center", vertical="center")

        def aplicar_estilo_tabela(ws, colunas_largura):
            for col in range(1, ws.max_column + 1):
                cell = ws.cell(1, col)
                cell.fill = cor_amarelo
                cell.font = fonte_negrito
                cell.alignment = alinhamento_centro

            ultima_linha = ws.max_row
            for col in range(1, ws.max_column + 1):
                cell = ws.cell(ultima_linha, col)
                cell.fill = cor_cinza
                cell.font = fonte_negrito

            for letra_col, largura in colunas_largura.items():
                ws.column_dimensions[letra_col].width = largura

        if (
            resumo_combustivel is not None
            and "Combustível (Vários itens)" in wb.sheetnames
        ):
            ws = wb["Combustível (Vários itens)"]
            aplicar_estilo_tabela(ws, {"A": 20, "B": 20, "C": 22, "D": 35, "E": 30})
            for row in range(2, ws.max_row + 1):
                ws.cell(row, 2).number_format = "#,##0.00"
                ws.cell(row, 3).number_format = "R$ #,##0.00"
                if ws.cell(row, 4).value:
                    ws.cell(row, 4).number_format = "#,##0"
                if ws.cell(row, 5).value:
                    ws.cell(row, 5).number_format = "#,##0"

        if resumo_arla is not None and "Arla 32" in wb.sheetnames:
            ws = wb["Arla 32"]
            aplicar_estilo_tabela(ws, {"A": 20, "B": 22})
            for row in range(2, ws.max_row + 1):
                ws.cell(row, 2).number_format = "R$ #,##0.00"

        ws = wb["Resumo por Posto"]
        aplicar_estilo_tabela(ws, {"A": 60, "B": 20, "C": 22})
        for row in range(2, ws.max_row + 1):
            ws.cell(row, 2).number_format = "#,##0.00"
            ws.cell(row, 3).number_format = "R$ #,##0.00"

        wb.save(caminho_saida)
        logger.info(f"Excel gerado com sucesso: {caminho_saida}")
        return True

    except Exception as e:
        logger.error(f"Falha ao gerar Excel para filial {nome_filial}: {e}")
        return False
