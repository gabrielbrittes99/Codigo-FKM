"""
Gerador de Relatórios de Frota por Filial
Conecta ao Bluefleet, extrai veículos ativos e gera as planilhas por filial.
"""

import os
import sys

import pandas as pd
import pyodbc
from dotenv import load_dotenv
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from src import config
from src.filial_mapping import normalizar_filial, EXCECOES_FORCADAS, PLACAS_EXCLUIDAS

load_dotenv()


def obter_conexao_bluefleet():
    """Retorna uma conexão pyodbc com o SQL Server BlueFleet."""
    host = os.getenv("DB_HOST")
    db = os.getenv("DB_NAME")
    user = os.getenv("DB_USER")
    pwd = os.getenv("DB_PASSWORD")

    if not all([host, db, user, pwd]):
        raise EnvironmentError(
            "Variáveis de ambiente do BlueFleet incompletas. "
            "Configure DB_HOST, DB_NAME, DB_USER e DB_PASSWORD no arquivo .env"
        )

    conn_str = (
        f"DRIVER={{ODBC Driver 18 for SQL Server}};"
        f"SERVER={host},1433;"
        f"DATABASE={db};"
        f"UID={user};"
        f"PWD={pwd};"
        f"TrustServerCertificate=yes;"
        f"Connection Timeout=30;"
    )

    return pyodbc.connect(conn_str)


def extrair_frota_bluefleet():
    import calendar
    print(f"🔄 Conectando ao Bluefleet para extrair dados históricos da frota para {config.MES}/{config.ANO}...")
    conn = obter_conexao_bluefleet()
    
    # 1. Carregar todos os veículos ativos
    query_vei = """
    SELECT
        Placa,
        Modelo,
        montadora,
        FilialOperacional,
        SituacaoVeiculo
    FROM
        dbo.Veiculos
    WHERE
        SituacaoVeiculo <> 'Vendido';
    """
    df_vei = pd.read_sql(query_vei, conn)
    
    # 2. Carregar movimentações operacionais (transferências entre filiais)
    query_mov = """
    SELECT
        Data_da_movimentação,
        Placa,
        Unidade_de_Origem,
        Unidade_de_Destino
    FROM
        dbo.Movimentos
    WHERE
        Unidade_movimentada = 'OPERAÇÃO'
    ORDER BY
        Placa, Data_da_movimentação;
    """
    df_mov = pd.read_sql(query_mov, conn)
    conn.close()

    # Normalizar placas para correspondência exata
    df_vei["Placa_Clean"] = df_vei["Placa"].astype(str).str.replace("-", "", regex=False).str.strip().str.upper()
    df_mov["Placa_Clean"] = df_mov["Placa"].astype(str).str.replace("-", "", regex=False).str.strip().str.upper()
    df_mov["Data_da_movimentação"] = pd.to_datetime(df_mov["Data_da_movimentação"])
    
    # Determinar período do fechamento
    mes_num = int(config.obter_numero_mes())
    ano = int(config.ANO)
    _, ultimo_dia = calendar.monthrange(ano, mes_num)
    start_date = pd.Timestamp(year=ano, month=mes_num, day=1, hour=0, minute=0, second=0)
    end_date = pd.Timestamp(year=ano, month=mes_num, day=ultimo_dia, hour=23, minute=59, second=59)
    
    print(f"   Período considerado: {start_date} até {end_date}")
    
    # Mapear movimentos por placa
    movs_by_plate = {p: g.sort_values("Data_da_movimentação") for p, g in df_mov.groupby("Placa_Clean")}
    
    historical_allocations = []

    for _, vei in df_vei.iterrows():
        placa = vei["Placa_Clean"]
        modelo = vei["Modelo"]
        situacao = vei["SituacaoVeiculo"]
        current_filial = vei["FilialOperacional"]

        # Ignorar placas excluídas (sinistrados, indenizados, baixados)
        if placa in PLACAS_EXCLUIDAS:
            continue

        # PRIORIDADE 1: Verificar se a placa está nas exceções forçadas
        if placa in EXCECOES_FORCADAS:
            filial_forcada = EXCECOES_FORCADAS[placa]
            # Normalizar nome da filial forçada
            filial_forcada = normalizar_filial(filial_forcada)

            # Adicionar apenas uma vez para a filial forçada
            historical_allocations.append({
                "Placa": vei["Placa"],
                "Modelo": modelo,
                "Montadora": vei.get("montadora", ""),
                "FilialOperacional": filial_forcada,
                "SituacaoVeiculo": situacao,
                "Data_Transferencia": "Placa forçada (mapeamento manual)",
                "Placa_Clean": placa
            })
            continue  # Pula a lógica normal de movimentações

        # PRIORIDADE 2: Determinar filial(is) no período usando movimentações
        p_movs = movs_by_plate.get(placa)

        # Resultado: lista de (filial, data_transferencia) onde o veículo esteve no mês
        filiais_no_periodo = []

        if p_movs is None or p_movs.empty:
            # Sem movimentações de OPERAÇÃO no banco = veículo estável
            filiais_no_periodo.append((current_filial, None))
        else:
            # Separar movimentações por período
            movs_antes = p_movs[p_movs["Data_da_movimentação"] < start_date]
            movs_durante = p_movs[(p_movs["Data_da_movimentação"] >= start_date) & (p_movs["Data_da_movimentação"] <= end_date)]
            movs_depois = p_movs[p_movs["Data_da_movimentação"] > end_date]

            if not movs_antes.empty:
                # Tem movimentações ANTES do mês: veículo estava no destino da última
                filiais_no_periodo.append((
                    movs_antes.iloc[-1]["Unidade_de_Destino"],
                    movs_antes.iloc[-1]["Data_da_movimentação"]
                ))
            elif not movs_durante.empty:
                # Sem movimentações antes, mas tem durante: veículo veio da origem
                filiais_no_periodo.append((
                    movs_durante.iloc[0]["Unidade_de_Origem"],
                    None
                ))
            elif not movs_depois.empty:
                # Só tem movimentações DEPOIS do período
                # Veículo estava na ORIGEM da primeira movimentação futura
                filiais_no_periodo.append((
                    movs_depois.iloc[0]["Unidade_de_Origem"],
                    None
                ))

            # Adicionar destinos de movimentações DURANTE o mês
            if not movs_durante.empty:
                movs_durante_copy = movs_durante.copy()
                movs_durante_copy["Data_Dia"] = movs_durante_copy["Data_da_movimentação"].dt.date
                for _, group in movs_durante_copy.groupby("Data_Dia"):
                    group_sorted = group.sort_values("Data_da_movimentação")
                    filiais_no_periodo.append((
                        group_sorted.iloc[-1]["Unidade_de_Destino"],
                        group_sorted.iloc[-1]["Data_da_movimentação"]
                    ))

        # Adicionar alocações (deduplicar por filial)
        filiais_vistas = set()
        alguma_grit_encontrada = False
        for filial, data_transf in filiais_no_periodo:
            if pd.isna(filial):
                continue
            filial_str = str(filial)
            if filial_str in filiais_vistas:
                continue
            filiais_vistas.add(filial_str)

            if "GRIT" in filial_str.upper():
                alguma_grit_encontrada = True
                historical_allocations.append({
                    "Placa": vei["Placa"],
                    "Modelo": modelo,
                    "Montadora": vei.get("montadora", ""),
                    "FilialOperacional": filial_str,
                    "SituacaoVeiculo": situacao,
                    "Data_Transferencia": data_transf.strftime("%d/%m/%Y") if data_transf is not None else "Sem transferência no período",
                    "Placa_Clean": placa
                })

        # Fallback: se nenhuma filial GRITSCH foi encontrada pelas movimentações,
        # usar a FilialOperacional atual do cadastro de veículos
        if not alguma_grit_encontrada and "GRIT" in str(current_filial).upper():
            historical_allocations.append({
                "Placa": vei["Placa"],
                "Modelo": modelo,
                "Montadora": vei.get("montadora", ""),
                "FilialOperacional": current_filial,
                "SituacaoVeiculo": situacao,
                "Data_Transferencia": "Sem transferência no período",
                "Placa_Clean": placa
            })
                
    df_hist = pd.DataFrame(historical_allocations)
    
    # Remover duplicatas caso a normalização ou lógica gere registros idênticos
    if not df_hist.empty:
        df_hist = df_hist.drop_duplicates(subset=["Placa_Clean", "FilialOperacional"])
        df_hist["Placa"] = df_hist["Placa"].astype(str).str.replace("-", "", regex=False).str.strip().str.upper()
        df_hist = df_hist.drop(columns=["Placa_Clean"])
    
    print(f"✅ Extração concluída. Total de alocações de veículos no mês: {len(df_hist)}")
    return df_hist


def gerar_relatorio_frota(df_filial, filial, caminho_saida):
    """
    Gera o arquivo Excel com a frota da filial e aplica formatação.
    """
    with pd.ExcelWriter(caminho_saida, engine="openpyxl") as writer:
        df_filial.to_excel(writer, sheet_name="Relação de Veículos", index=False)

    # Aplicar formatação
    wb = load_workbook(caminho_saida)
    ws = wb["Relação de Veículos"]

    cor_amarelo = PatternFill(
        start_color="FFFF00", end_color="FFFF00", fill_type="solid"
    )
    cor_cinza = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
    fonte_negrito = Font(bold=True)
    alinhamento_centro = Alignment(horizontal="center", vertical="center")

    # Formatar cabeçalho
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(1, col)
        cell.fill = cor_amarelo
        cell.font = fonte_negrito
        cell.alignment = alinhamento_centro

    # Adicionar linha de totalização no final
    ultima_linha = ws.max_row + 1
    ws.cell(ultima_linha, 1).value = "Total de Veículos"
    ws.cell(ultima_linha, 1).fill = cor_cinza
    ws.cell(ultima_linha, 1).font = fonte_negrito

    ws.cell(ultima_linha, 2).value = len(df_filial)
    ws.cell(ultima_linha, 2).fill = cor_cinza
    ws.cell(ultima_linha, 2).font = fonte_negrito
    ws.cell(ultima_linha, 2).alignment = alinhamento_centro

    for col in range(3, ws.max_column + 1):
        ws.cell(ultima_linha, col).fill = cor_cinza

    # Ajustar larguras
    ws.column_dimensions["A"].width = 15  # Placa
    ws.column_dimensions["B"].width = 40  # Modelo
    ws.column_dimensions["C"].width = 20  # Montadora
    ws.column_dimensions["D"].width = 30  # FilialOperacional
    ws.column_dimensions["E"].width = 25  # SituacaoVeiculo
    ws.column_dimensions["F"].width = 30  # Data_Transferencia

    wb.save(caminho_saida)


def main():
    print("=" * 80)
    print(f"GERADOR DE RESUMOS DE FROTA - {config.MES}/{config.ANO}")
    print("=" * 80)

    try:
        df = extrair_frota_bluefleet()
    except Exception as e:
        print(f"\n❌ ERRO ao conectar no banco de dados: {e}")
        sys.exit(1)

    if df.empty:
        print("⚠️  Nenhum veículo encontrado com os critérios fornecidos.")
        sys.exit(0)

    # Normalizar nomes das filiais para ficar igual Manutenção e Combustível
    df["Filial_Normalizada"] = df["FilialOperacional"].apply(normalizar_filial)

    # Se normalizar_filial não mapear algumas, garantimos que fiquem parecidas com o padrão
    # Mas a normalização lida bem com 'GRITSCH - XXX'

    filiais_unicas = df["Filial_Normalizada"].dropna().unique()

    print(f"\n✅ {len(filiais_unicas)} filiais únicas encontradas")
    print("\n" + "=" * 80)
    print("GERANDO ARQUIVOS POR FILIAL")
    print("=" * 80)

    for idx, filial in enumerate(filiais_unicas, 1):
        # Filtramos pelas filiais normalizadas para o Excel
        df_filial = df[df["Filial_Normalizada"] == filial].copy()

        # Removemos a coluna auxiliar antes de salvar
        df_filial = df_filial.drop(columns=["Filial_Normalizada"])

        # Criar pasta da filial
        pasta_filial = config.obter_caminho_saida_filial(filial)
        os.makedirs(pasta_filial, exist_ok=True)

        nome_filial_limpo = "".join(
            c for c in str(filial) if c.isalnum() or c in (" ", "_")
        ).rstrip()
        nome_arquivo_saida = f"Frota - {nome_filial_limpo}.xlsx"
        caminho_completo_saida = os.path.join(pasta_filial, nome_arquivo_saida)

        print(f"\n[{idx}/{len(filiais_unicas)}] 🚙 {filial}")
        print(f"      Veículos: {len(df_filial)}")
        print(f"      Pasta: {filial}/")
        print(f"      Gerando: {nome_arquivo_saida}")

        if os.path.exists(caminho_completo_saida):
            try:
                os.remove(caminho_completo_saida)
                print(f"      🗑️  Arquivo antigo removido")
            except Exception as e:
                print(f"      ⚠️  Não foi possível remover: {e}")
                continue

        gerar_relatorio_frota(df_filial, filial, caminho_completo_saida)
        print(f"      ✅ Concluído!")

    print("\n" + "=" * 80)
    print("SALVANDO ARQUIVO GERAL")
    print("=" * 80)

    mes_num = config.obter_numero_mes()
    ano_short = config.ANO[-2:]
    nome_arquivo_geral = f"{mes_num}{ano_short} FROTA GRITSCH TRANSPORTES GERAL.xlsx"
    pasta_periodo = os.path.join(config.DIRETORIO_BASE_SAIDA, config.PASTA_PERIODO)
    os.makedirs(pasta_periodo, exist_ok=True)
    caminho_completo_geral = os.path.join(pasta_periodo, nome_arquivo_geral)

    print(f"\n🔄 Salvando arquivo geral...")
    df.drop(columns=["Filial_Normalizada"]).to_excel(
        caminho_completo_geral, index=False, engine="openpyxl"
    )
    print(f"✅ Arquivo geral salvo!")

    print("\n" + "=" * 80)
    print("🎉 PROCESSO CONCLUÍDO COM SUCESSO!")
    print("=" * 80)
    print(f"\n📁 Arquivos salvos em: {pasta_periodo}")


if __name__ == "__main__":
    main()
