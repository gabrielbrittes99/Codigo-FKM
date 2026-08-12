"""
Gera o FKM da GRITSCH - MATRIZ.

A MATRIZ é um hub de trânsito com três tipos de veículos:
  1. Veículos de filiais em manutenção pesada (ficam 3 semanas a 1 mês)
  2. Veículos novos em preparação para envio às filiais (destino pode ser "A definir")
  3. Veículos vindos de REFERÊNCIA sendo incorporados à Gritsch

Colunas especiais (diferente das filiais normais):
  - Filial Origem: de onde o veículo veio
  - Data Chegada: quando chegou na MATRIZ
  - Situação: Em manutenção / Em preparação / Transferido para [filial]
  - Destino: filial de destino ou "A definir"
"""

import calendar
import os

import pandas as pd
import pyodbc
from dotenv import load_dotenv
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src import config
from src.filial_mapping import normalizar_filial, EXCECOES_FORCADAS

load_dotenv()


def _obter_conexao():
    host = os.getenv("DB_HOST")
    db = os.getenv("DB_NAME")
    user = os.getenv("DB_USER")
    pwd = os.getenv("DB_PASSWORD")
    conn_str = (
        f"DRIVER={{ODBC Driver 18 for SQL Server}};"
        f"SERVER={host},1433;DATABASE={db};UID={user};PWD={pwd};"
        f"TrustServerCertificate=yes;Connection Timeout=30;"
    )
    return pyodbc.connect(conn_str)


def _classificar_situacao(origem, destino_no_mes):
    """
    Classifica a situação do veículo na MATRIZ com base na origem.
    Se já saiu no mês, indica para onde foi.
    """
    if destino_no_mes:
        return f"Transferido para {destino_no_mes}"

    origem_upper = str(origem).upper() if origem else ""
    if "REFERÊNCIA" in origem_upper or "REFERENCIA" in origem_upper:
        return "Em incorporação"
    if "GRITSCH" in origem_upper:
        return "Em manutenção"
    return "Em preparação"


def carregar_dados_movimentos_matriz(mes_num, ano):
    """
    Busca no banco todos os veículos que passaram pela MATRIZ no período,
    com origem e destino de cada um.
    """
    conn = _obter_conexao()

    _, ultimo_dia = calendar.monthrange(ano, mes_num)
    inicio = pd.Timestamp(year=ano, month=mes_num, day=1)
    fim = pd.Timestamp(year=ano, month=mes_num, day=ultimo_dia, hour=23, minute=59)

    query_mov = """
    SELECT
        Data_da_movimentação,
        Placa,
        Unidade_de_Origem,
        Unidade_de_Destino
    FROM dbo.Movimentos
    WHERE Unidade_movimentada = 'OPERAÇÃO'
    ORDER BY Placa, Data_da_movimentação
    """
    df_mov = pd.read_sql(query_mov, conn)

    # Busca todos os veículos (inclusive vendidos) para lookup de modelo
    query_vei_todos = "SELECT Placa, Modelo, FilialOperacional, SituacaoVeiculo FROM dbo.Veiculos"
    df_vei_todos = pd.read_sql(query_vei_todos, conn)
    conn.close()

    df_vei_todos["Placa_Clean"] = (
        df_vei_todos["Placa"].astype(str).str.replace("-", "").str.strip().str.upper()
    )
    df_mov["Placa_Clean"] = (
        df_mov["Placa"].astype(str).str.replace("-", "").str.strip().str.upper()
    )

    # Índice para lookup de modelo (usa todos, incluindo vendidos)
    modelo_lookup = df_vei_todos.set_index("Placa_Clean")["Modelo"].to_dict()

    # Somente veículos ativos — vendidos saem da frota, não entram no FKM da MATRIZ
    # Comparação case-insensitive: banco armazena como "VENDIDO" (maiúsculas)
    placas_ativas = set(
        df_vei_todos[df_vei_todos["SituacaoVeiculo"].str.upper() != "VENDIDO"]["Placa_Clean"]
    )

    df_mov["Data_da_movimentação"] = pd.to_datetime(df_mov["Data_da_movimentação"])

    MATRIZ_NAMES = {"GRITSCH - MATRIZ", "GRITSCH"}

    registros = []

    for placa, movs in df_mov.groupby("Placa_Clean"):
        # Excluir veículos vendidos — eles saem da frota ativa, não pertencem ao FKM
        if placa not in placas_ativas:
            continue

        movs = movs.sort_values("Data_da_movimentação")

        # Movimentos que entram NA MATRIZ (destino = MATRIZ)
        entradas = movs[movs["Unidade_de_Destino"].isin(MATRIZ_NAMES)]
        # Movimentos que saem DA MATRIZ (origem = MATRIZ)
        saidas = movs[movs["Unidade_de_Origem"].isin(MATRIZ_NAMES)]

        # Determinar se o veículo estava na MATRIZ durante o período
        # 1) Chegou na MATRIZ antes ou durante o período
        entradas_ate_fim = entradas[entradas["Data_da_movimentação"] <= fim]
        if entradas_ate_fim.empty:
            # Sem entrada registrada; verificar se FilialOperacional atual é MATRIZ
            vei_info = df_vei_todos[df_vei_todos["Placa_Clean"] == placa]
            if vei_info.empty:
                continue
            filial_atual = str(vei_info.iloc[0]["FilialOperacional"])
            if filial_atual not in MATRIZ_NAMES:
                continue
            # Veículo está na MATRIZ mas sem movimento de entrada registrado
            data_chegada = None
            origem = "Não registrada"
        else:
            ultima_entrada = entradas_ate_fim.iloc[-1]
            data_chegada = ultima_entrada["Data_da_movimentação"]
            origem = ultima_entrada["Unidade_de_Origem"]

            # 2) Verificar se saiu da MATRIZ ANTES do início do período
            saidas_antes = saidas[saidas["Data_da_movimentação"] < inicio]
            if not saidas_antes.empty:
                ultima_saida = saidas_antes.iloc[-1]
                # Saiu e pode ter voltado? Só inclui se a última entrada foi DEPOIS da última saída
                if data_chegada is not None and ultima_saida["Data_da_movimentação"] >= data_chegada:
                    continue  # saiu e não voltou antes do período

        # 3) Verificar se saiu DA MATRIZ durante o período
        saidas_no_mes = saidas[
            (saidas["Data_da_movimentação"] >= inicio) &
            (saidas["Data_da_movimentação"] <= fim)
        ]
        destino_no_mes = None
        data_saida = None
        if not saidas_no_mes.empty:
            primeira_saida = saidas_no_mes.iloc[0]
            destino_no_mes = primeira_saida["Unidade_de_Destino"]
            data_saida = primeira_saida["Data_da_movimentação"]

        modelo = modelo_lookup.get(placa, "")

        registros.append({
            "Placa_Clean": placa,
            "Modelo": modelo,
            "Filial_Origem": normalizar_filial(origem) if origem != "Não registrada" else origem,
            "Data_Chegada": data_chegada.strftime("%d/%m/%Y") if data_chegada else "Não registrada",
            "Destino": normalizar_filial(destino_no_mes) if destino_no_mes else "A definir",
            "Data_Saida": data_saida.strftime("%d/%m/%Y") if data_saida else "",
            "Situacao": _classificar_situacao(origem, normalizar_filial(destino_no_mes) if destino_no_mes else None),
        })

    return pd.DataFrame(registros)


def carregar_combustivel_matriz(pasta_periodo):
    """Lê o arquivo de combustível gerado pelo pipeline para MATRIZ."""
    caminho = os.path.join(pasta_periodo, "GRITSCH - MATRIZ", "Combustivel - GRITSCH  MATRIZ.xlsx")
    if not os.path.exists(caminho):
        # Tentar variações de nome
        for nome in os.listdir(os.path.join(pasta_periodo, "GRITSCH - MATRIZ") if os.path.exists(os.path.join(pasta_periodo, "GRITSCH - MATRIZ")) else pasta_periodo):
            if "Combustivel" in nome and "MATRIZ" in nome:
                caminho = os.path.join(pasta_periodo, "GRITSCH - MATRIZ", nome)
                break
        else:
            return pd.DataFrame()

    df = pd.read_excel(caminho)
    df["Placa_Clean"] = df["Placa"].astype(str).str.replace("-", "").str.strip().str.upper()

    resumo = df.groupby("Placa_Clean").agg(
        Litros=("Litragem", "sum"),
        Valor_Comb=("Valor total", "sum"),
    ).reset_index()

    if "Combustivel" in df.columns:
        arla = df[df["Combustivel"].str.upper().str.contains("ARLA", na=False)].groupby("Placa_Clean")["Valor total"].sum()
        resumo["Arla"] = resumo["Placa_Clean"].map(arla).fillna(0)
    else:
        resumo["Arla"] = 0

    return resumo


def carregar_manutencao_matriz(pasta_periodo):
    """Lê o arquivo de manutenção gerado pelo pipeline para MATRIZ."""
    pasta_manut = os.path.join(pasta_periodo, "GRITSCH - MATRIZ")
    if not os.path.exists(pasta_manut):
        return pd.DataFrame()

    for nome in os.listdir(pasta_manut):
        if "Manutencao" in nome or "Manutenção" in nome:
            caminho = os.path.join(pasta_manut, nome)
            df = pd.read_excel(caminho, sheet_name="Dados Brutos")
            df["Placa_Clean"] = df["Placa"].astype(str).str.replace("-", "").str.strip().str.upper()

            col_natureza = "Natureza_Correta" if "Natureza_Correta" in df.columns else "TipoItem"
            resumo = df.pivot_table(
                index="Placa_Clean",
                columns=col_natureza,
                values="ValorTotal",
                aggfunc="sum",
                fill_value=0,
            ).reset_index()
            return resumo

    return pd.DataFrame()


def gerar_fkm_matriz(caminho_saida, df_movs, df_comb, df_manut):
    """Monta e salva o Excel do FKM MATRIZ."""
    # Merge principal
    df = df_movs.copy()

    if not df_comb.empty:
        df = df.merge(df_comb, on="Placa_Clean", how="left")
    else:
        df["Litros"] = 0
        df["Valor_Comb"] = 0
        df["Arla"] = 0

    if not df_manut.empty:
        df = df.merge(df_manut, on="Placa_Clean", how="left")

    df = df.fillna(0)

    # Ordenar por situação e origem
    ordem_situacao = {"Em manutenção": 0, "Em incorporação": 1, "Em preparação": 2}
    df["_ordem"] = df["Situacao"].map(lambda s: 0 if s.startswith("Transferido") else ordem_situacao.get(s, 3))
    df = df.sort_values(["_ordem", "Filial_Origem", "Placa_Clean"]).drop(columns=["_ordem"])

    # Colunas fixas que sempre aparecem
    cols_fixas = ["Placa_Clean", "Modelo", "Filial_Origem", "Data_Chegada",
                  "Situacao", "Destino", "Data_Saida",
                  "Litros", "Valor_Comb", "Arla"]

    # Colunas de manutenção por natureza (dinâmicas)
    cols_manut = [c for c in df.columns if c not in cols_fixas and c != "Placa_Clean"]

    colunas_finais = cols_fixas + cols_manut
    colunas_finais = [c for c in colunas_finais if c in df.columns]

    df_export = df[colunas_finais].rename(columns={
        "Placa_Clean": "Placa",
        "Filial_Origem": "Filial de Origem",
        "Data_Chegada": "Data Chegada",
        "Data_Saida": "Data Saída",
        "Situacao": "Situação",
        "Litros": "Litros Comb.",
        "Valor_Comb": "Valor Comb. (R$)",
        "Arla": "Arla (R$)",
    })

    with pd.ExcelWriter(caminho_saida, engine="openpyxl") as writer:
        df_export.to_excel(writer, sheet_name="FKM MATRIZ", index=False)

    # Formatação
    wb = load_workbook(caminho_saida)
    ws = wb["FKM MATRIZ"]

    cor_azul = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
    cor_amarelo = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    cor_cinza = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
    fonte_cab = Font(bold=True, color="FFFFFF")
    fonte_total = Font(bold=True)
    alinhamento_centro = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Cabeçalho
    ws.insert_rows(1)
    ws.insert_rows(1)
    ws["A1"] = f"FKM — GRITSCH MATRIZ  |  {config.MES} {config.ANO}"
    ws["A1"].font = Font(bold=True, size=14, color="1F3864")
    ws["A2"] = "Veículos em manutenção pesada, preparação e incorporação — preenchimento: equipe de manutenção"
    ws["A2"].font = Font(italic=True, size=10)

    linha_cab = 3
    for cell in ws[linha_cab]:
        if cell.value:
            cell.fill = cor_azul
            cell.font = fonte_cab
            cell.alignment = alinhamento_centro

    # Destacar veículos sem destino definido ("A definir")
    col_destino = None
    for cell in ws[linha_cab]:
        if cell.value == "Destino":
            col_destino = cell.column
            break

    if col_destino:
        for row in ws.iter_rows(min_row=linha_cab + 1, max_row=ws.max_row):
            if row[col_destino - 1].value == "A definir":
                row[col_destino - 1].fill = cor_amarelo

    # Linha de totais
    ultima = ws.max_row + 1
    ws.cell(ultima, 1).value = f"TOTAL  ({len(df_export)} veículos)"
    ws.cell(ultima, 1).font = fonte_total

    # Somar colunas numéricas
    for col_idx, cell in enumerate(ws[linha_cab], 1):
        if cell.value in ("Litros Comb.", "Valor Comb. (R$)", "Arla (R$)") or (
            cell.value and any(nat in str(cell.value) for nat in ["Peças", "Mão", "Pneu", "Lataria", "Serviço"])
        ):
            total = sum(
                (ws.cell(r, col_idx).value or 0)
                for r in range(linha_cab + 1, ws.max_row)
                if isinstance(ws.cell(r, col_idx).value, (int, float))
            )
            ws.cell(ultima, col_idx).value = round(total, 2)
            ws.cell(ultima, col_idx).number_format = 'R$ #,##0.00'
            ws.cell(ultima, col_idx).font = fonte_total

    for cell in ws[ultima]:
        if cell.value is not None:
            cell.fill = cor_cinza

    # Larguras
    ws.column_dimensions["A"].width = 12   # Placa
    ws.column_dimensions["B"].width = 32   # Modelo
    ws.column_dimensions["C"].width = 25   # Filial de Origem
    ws.column_dimensions["D"].width = 14   # Data Chegada
    ws.column_dimensions["E"].width = 28   # Situação
    ws.column_dimensions["F"].width = 25   # Destino
    ws.column_dimensions["G"].width = 12   # Data Saída
    ws.column_dimensions["H"].width = 12   # Litros
    ws.column_dimensions["I"].width = 16   # Valor Comb.
    ws.column_dimensions["J"].width = 12   # Arla

    ws.row_dimensions[linha_cab].height = 35

    wb.save(caminho_saida)


def main():
    print("=" * 70)
    print(f"FKM GRITSCH MATRIZ — {config.MES}/{config.ANO}")
    print("=" * 70)

    mes_num = int(config.obter_numero_mes())
    ano = int(config.ANO)
    pasta_periodo = os.path.join(config.DIRETORIO_BASE_SAIDA, config.PASTA_PERIODO)

    print("\n🔍 Buscando veículos na MATRIZ via tabela de movimentos...")
    df_movs = carregar_dados_movimentos_matriz(mes_num, ano)

    if df_movs.empty:
        print("⚠️  Nenhum veículo encontrado na MATRIZ no período.")
        return

    print(f"   ✅ {len(df_movs)} veículos identificados na MATRIZ")
    for sit, grp in df_movs.groupby("Situacao"):
        print(f"      - {sit}: {len(grp)} veículos")

    print("\n📊 Carregando dados de combustível...")
    df_comb = carregar_combustivel_matriz(pasta_periodo)
    print(f"   {'✅' if not df_comb.empty else '⚠️ '} {len(df_comb)} placas com combustível")

    print("\n🔧 Carregando dados de manutenção...")
    df_manut = carregar_manutencao_matriz(pasta_periodo)
    print(f"   {'✅' if not df_manut.empty else '⚠️ '} {len(df_manut)} placas com manutenção")

    # Salvar na pasta da MATRIZ
    pasta_matriz = os.path.join(pasta_periodo, "GRITSCH - MATRIZ")
    os.makedirs(pasta_matriz, exist_ok=True)

    mes_num_str = config.obter_numero_mes()
    ano_short = config.ANO[-2:]
    nome_arquivo = f"FKM MATRIZ {mes_num_str}{ano_short}.xlsx"
    caminho_saida = os.path.join(pasta_matriz, nome_arquivo)

    print(f"\n📝 Gerando FKM MATRIZ...")
    gerar_fkm_matriz(caminho_saida, df_movs, df_comb, df_manut)

    print(f"\n✅ FKM salvo em: {caminho_saida}")

    a_definir = (df_movs["Destino"] == "A definir").sum()
    if a_definir:
        print(f"\n⚠️  {a_definir} veículo(s) com destino 'A definir' — marcados em amarelo no arquivo.")

    print("=" * 70)


if __name__ == "__main__":
    main()
