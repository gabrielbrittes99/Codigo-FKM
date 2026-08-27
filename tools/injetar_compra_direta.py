"""
Injeta a compra direta de combustível (fora da TruckPag) no arquivo oficial de
cada filial, como uma aba extra.

Por que isto existe: o arquivo "Combustivel - GRITSCH X.xlsx" que mandamos para
a filial preencher o FKM hoje só traz o que passou pela TruckPag. Quando a
filial compra diesel ou Arla direto do posto (fora da rede credenciada), essa
nota nunca aparece nesse arquivo — então a filial não tem como saber que
precisa declarar aquele valor no FKM, e o FKM nunca fecha com a realidade.

O que este script faz: lê os lançamentos de compra direta do financeiro
(mesma fonte de `src.torre_dados._extrair_combustivel_fora`), agrupa por
filial e escreve, no arquivo oficial já gerado pelo fechamento, uma aba
"Compra Direta (Fora TruckPag)" com o detalhe (data, fornecedor, nota,
natureza, valor) e uma aba "Resumo Geral" com TruckPag + Direta = Total Real.

Uso:
    python -m tools.injetar_compra_direta --mes 7 --ano 2026
    python -m tools.injetar_compra_direta --mes 7 --ano 2026 --sem-cache
"""

import argparse
import glob
import os
import pickle
import unicodedata

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src import config
from src.torre_dados import (
    DIR_CACHE,
    FILIAIS_NAO_OPERACIONAIS,
    ROTULO_NAO_ALOCADO,
    ROTULO_SEM_CENTRO,
    ROTULO_SEM_EQUIVALENCIA,
    _extrair_combustivel_fora,
    centro_custo_para_filial,
)

NOME_ABA_DETALHE = "Compra Direta (Fora TruckPag)"
NOME_ABA_RESUMO = "Resumo Geral"

COR_AMARELO = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
COR_CINZA = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
COR_ALERTA = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
FONTE_NEGRITO = Font(bold=True)
FONTE_NEGRITO_GRANDE = Font(bold=True, size=13)
ALINHAMENTO_CENTRO = Alignment(horizontal="center", vertical="center")


def _sem_acento(texto):
    return "".join(
        c for c in unicodedata.normalize("NFD", str(texto))
        if unicodedata.category(c) != "Mn"
    ).upper().strip()


def _normalizar(texto):
    """Remove tudo que não é letra/número, para comparar nomes com folga."""
    return "".join(c for c in _sem_acento(texto) if c.isalnum())


def carregar_compra_direta(mes, ano, usar_cache=True):
    """Lançamentos de compra direta do mês, de cache ou do banco."""
    caminho_cache = os.path.join(DIR_CACHE, f"torre_bruto_mensal_{ano}_{mes:02d}.pkl")

    if usar_cache and os.path.exists(caminho_cache):
        with open(caminho_cache, "rb") as arquivo:
            bruto = pickle.load(arquivo)
        df = bruto.get("combustivel_fora")
        if df is not None and not df.empty:
            print(f"💾 Compra direta lida do cache ({caminho_cache}).")
            return df[df["mes"] == mes].copy()

    print("📊 Cache não tem compra direta deste mês — consultando o banco...")
    return _extrair_combustivel_fora([mes], ano)


def localizar_pastas_filiais(mes_nome, ano):
    """{código da filial: caminho da pasta}, varrendo o que existe de verdade."""
    pasta_periodo = os.path.join(config.DIRETORIO_BASE_SAIDA, f"{mes_nome} {ano}")
    if not os.path.isdir(pasta_periodo):
        raise FileNotFoundError(
            f"Pasta do fechamento não encontrada: {pasta_periodo}\n"
            f"Rode o fechamento do mês antes (fechar_mes.py)."
        )

    pastas = {}
    for nome in os.listdir(pasta_periodo):
        caminho = os.path.join(pasta_periodo, nome)
        if not os.path.isdir(caminho):
            continue
        # "GRITSCH - CWB BASE" -> "CWBBASE" para comparar com o código da filial
        codigo_normalizado = _normalizar(nome.replace("GRITSCH", "", 1))
        pastas[codigo_normalizado] = caminho
    return pastas


def localizar_arquivo_combustivel(pasta_filial):
    candidatos = glob.glob(os.path.join(pasta_filial, "Combustivel - *.xlsx"))
    return candidatos[0] if candidatos else None


def nome_arquivo_novo(pasta_filial, codigo):
    """Nome no mesmo padrão de executar_resumos.py, para a filial que nunca
    teve arquivo de combustível (só abastece fora da TruckPag — caso do ITR)."""
    nome_pasta = os.path.basename(pasta_filial)  # "GRITSCH - ITR"
    nome_limpo = "".join(c for c in nome_pasta if c.isalnum() or c in (" ", "_")).rstrip()
    return os.path.join(pasta_filial, f"Combustivel - {nome_limpo}.xlsx")


def _formatar_moeda(ws, coluna, linha_inicio, linha_fim):
    for linha in range(linha_inicio, linha_fim + 1):
        ws.cell(linha, coluna).number_format = "R$ #,##0.00"


def _ajustar_larguras(ws, larguras):
    for indice, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(indice)].width = largura


def escrever_aba_detalhe(wb, lancamentos, total_truckpag):
    if NOME_ABA_DETALHE in wb.sheetnames:
        del wb[NOME_ABA_DETALHE]
    ws = wb.create_sheet(NOME_ABA_DETALHE, 0)

    colunas = ["Data", "Fornecedor", "Nº Documento", "Natureza", "Valor", "Descrição"]
    ws.append(colunas)
    for col in range(1, len(colunas) + 1):
        celula = ws.cell(1, col)
        celula.fill = COR_AMARELO
        celula.font = FONTE_NEGRITO
        celula.alignment = ALINHAMENTO_CENTRO

    linha = 2
    for _, r in lancamentos.sort_values("data").iterrows():
        ws.append([
            r["data"].strftime("%d/%m/%Y") if pd.notna(r["data"]) else "",
            str(r.get("fornecedor") or ""),
            str(r.get("numero_documento") or ""),
            str(r.get("natureza") or ""),
            float(r["valor"]),
            str(r.get("descricao") or ""),
        ])
        linha += 1

    total_direta = float(lancamentos["valor"].sum())
    ws.append(["", "", "", "TOTAL COMPRA DIRETA", total_direta, ""])
    for col in range(1, len(colunas) + 1):
        celula = ws.cell(linha, col)
        celula.fill = COR_CINZA
        celula.font = FONTE_NEGRITO

    _formatar_moeda(ws, 5, 2, linha)
    _ajustar_larguras(ws, [12, 34, 16, 26, 14, 46])

    aviso = ws.cell(linha + 2, 1)
    aviso.value = (
        "⚠ Este valor NÃO está incluído nas abas de TruckPag deste arquivo. "
        "Ele precisa ser somado ao preencher o FKM do mês."
    )
    aviso.font = Font(bold=True, color="9C0006")
    ws.merge_cells(start_row=linha + 2, start_column=1, end_row=linha + 2, end_column=6)

    return total_direta


def escrever_aba_resumo(wb, total_truckpag, total_direta):
    if NOME_ABA_RESUMO in wb.sheetnames:
        del wb[NOME_ABA_RESUMO]
    ws = wb.create_sheet(NOME_ABA_RESUMO, 0)

    ws["A1"] = "Resumo do combustível do mês"
    ws["A1"].font = FONTE_NEGRITO_GRANDE
    ws.merge_cells("A1:B1")

    linhas = [
        ("Via TruckPag (rede credenciada)", total_truckpag),
        ("Compra direta (fora da rede)", total_direta),
        ("TOTAL REAL DO MÊS", total_truckpag + total_direta),
    ]
    for i, (rotulo, valor) in enumerate(linhas, start=3):
        ws.cell(i, 1, rotulo)
        c_valor = ws.cell(i, 2, valor)
        c_valor.number_format = "R$ #,##0.00"
        if "TOTAL" in rotulo:
            ws.cell(i, 1).font = FONTE_NEGRITO
            c_valor.font = FONTE_NEGRITO
            ws.cell(i, 1).fill = COR_CINZA
            c_valor.fill = COR_CINZA

    ws.cell(7, 1, "O FKM deste mês deve fechar com o TOTAL REAL, não só com a TruckPag.")
    ws.cell(7, 1).font = Font(italic=True, color="9C0006")
    ws.merge_cells("A7:D7")

    _ajustar_larguras(ws, [34, 16])
    return ws


def somar_truckpag_do_arquivo(caminho_arquivo):
    """Lê o total já presente na aba 'Combustível (Vários itens)' do arquivo oficial."""
    try:
        df = pd.read_excel(caminho_arquivo, sheet_name="Combustível (Vários itens)")
    except ValueError:
        return 0.0
    linha_total = df[df.iloc[:, 0].astype(str).str.contains("Total Geral", na=False)]
    if linha_total.empty:
        return 0.0
    coluna_valor = next((c for c in df.columns if "valor" in c.lower()), None)
    return float(linha_total[coluna_valor].iloc[0]) if coluna_valor else 0.0


def processar(mes, ano, usar_cache=True):
    meses_nomes = {1: "Janeiro", 2: "Fevereiro", 3: "Março", 4: "Abril", 5: "Maio",
                   6: "Junho", 7: "Julho", 8: "Agosto", 9: "Setembro", 10: "Outubro",
                   11: "Novembro", 12: "Dezembro"}
    mes_nome = meses_nomes[mes]

    fora = carregar_compra_direta(mes, ano, usar_cache)
    if fora.empty:
        print("✅ Nenhum lançamento de compra direta neste mês. Nada a fazer.")
        return

    fora = fora.copy()
    fora["cod"] = fora["filial"].apply(centro_custo_para_filial)

    pastas = localizar_pastas_filiais(mes_nome, ano)

    print("=" * 78)
    print(f"COMPRA DIRETA — {mes_nome}/{ano}")
    print("=" * 78)
    print(f"Total de lançamentos: {len(fora)}  |  Valor: R$ {fora['valor'].sum():,.2f}\n")

    processadas, sem_pasta, nao_operacionais = [], [], []

    for codigo, grupo in fora.groupby("cod"):
        valor_codigo = float(grupo["valor"].sum())

        if codigo in (ROTULO_NAO_ALOCADO, ROTULO_SEM_CENTRO, ROTULO_SEM_EQUIVALENCIA):
            sem_pasta.append((codigo, valor_codigo, len(grupo)))
            continue
        if codigo in FILIAIS_NAO_OPERACIONAIS:
            nao_operacionais.append((codigo, valor_codigo, len(grupo)))
            continue

        pasta = pastas.get(_normalizar(codigo))
        if not pasta:
            sem_pasta.append((codigo, valor_codigo, len(grupo)))
            continue

        arquivo = localizar_arquivo_combustivel(pasta)
        criado_do_zero = arquivo is None
        if criado_do_zero:
            # Filial sem nenhuma transação TruckPag no mês (caso do ITR): o
            # arquivo nunca foi gerado pelo fechamento normal. Cria-se aqui só
            # com o que existe de verdade — a compra direta.
            arquivo = nome_arquivo_novo(pasta, codigo)
            wb = Workbook()
            wb.remove(wb.active)
            total_truckpag = 0.0
        else:
            total_truckpag = somar_truckpag_do_arquivo(arquivo)
            wb = load_workbook(arquivo)

        total_direta = escrever_aba_detalhe(wb, grupo, total_truckpag)
        escrever_aba_resumo(wb, total_truckpag, total_direta)
        wb.active = wb.sheetnames.index(NOME_ABA_RESUMO)
        wb.save(arquivo)

        processadas.append((codigo, valor_codigo, len(grupo), arquivo))
        marca = " [ARQUIVO NOVO — filial sem TruckPag]" if criado_do_zero else ""
        print(f"✅ {codigo:<20} R$ {valor_codigo:>11,.2f}  ({len(grupo)} lançamentos)  "
              f"-> {os.path.basename(arquivo)}{marca}")

    if nao_operacionais:
        print("\nℹ️  Não operacionais (fora do relatório executivo, não injetadas):")
        for codigo, valor, n in nao_operacionais:
            print(f"   {codigo:<20} R$ {valor:>11,.2f}  ({n})")

    if sem_pasta:
        print("\n⚠️  SEM PASTA CORRESPONDENTE — precisa de atenção manual:")
        for codigo, valor, n in sem_pasta:
            print(f"   {codigo:<20} R$ {valor:>11,.2f}  ({n})")

    print("\n" + "=" * 78)
    print(f"RESUMO: {len(processadas)} filiais atualizadas | "
          f"{len(sem_pasta)} sem pasta | {len(nao_operacionais)} não operacionais")
    print("=" * 78)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="Injeta a compra direta de combustível nos arquivos oficiais do fechamento.")
    parser.add_argument("--mes", type=int, required=True, help="Mês do fechamento (1-12).")
    parser.add_argument("--ano", type=int, required=True)
    parser.add_argument("--sem-cache", action="store_true",
                        help="Reconsulta o banco em vez de usar o cache salvo.")
    args = parser.parse_args()

    processar(args.mes, args.ano, usar_cache=not args.sem_cache)
