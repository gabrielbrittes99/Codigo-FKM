"""
Gera relatório informativo de custos fora do escopo FKM.

Inclui categorias que são rastreadas mas não fazem parte do fechamento
mensal das filiais operacionais:
  - Veículos para Venda: custos de preparação de veículos em processo de venda
  - Referência: veículos de outras empresas do grupo que usam o mesmo sistema

O relatório é salvo em Dados Tratados/{PERIODO}/ junto com os demais arquivos.
"""

import os

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src import config


def _estilo_cabecalho(ws, linha, cor_hex):
    fill = PatternFill(start_color=cor_hex, end_color=cor_hex, fill_type="solid")
    font = Font(bold=True, color="FFFFFF")
    alinhamento = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for cell in ws[linha]:
        if cell.value is not None:
            cell.fill = fill
            cell.font = font
            cell.alignment = alinhamento


def _estilo_total(ws, linha):
    fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
    font = Font(bold=True)
    for cell in ws[linha]:
        if cell.value is not None:
            cell.fill = fill
            cell.font = font


def _ajustar_colunas(ws, larguras):
    for col_letra, largura in larguras.items():
        ws.column_dimensions[col_letra].width = largura


def _eh_referencia(filial):
    f = str(filial).upper()
    return f.startswith("REFERÊNCIA") or f.startswith("REFERENCIA")


def _gerar_aba_veiculos_venda(wb, df_comb, df_manut):
    """Aba com custos de veículos em processo de venda."""
    ws = wb.create_sheet("Veículos para Venda")

    mask_comb = df_comb["Filial_Final"].str.upper().str.contains("VEÍCULOS PARA VENDA|VEICULOS PARA VENDA", na=False)
    df_v = df_comb[mask_comb].copy()

    # Manutenção: cruzar pelas placas que estão em veículos para venda
    placas_venda = set(df_v["Placa_Clean"].dropna().unique())
    df_m = df_manut[df_manut["Placa_Clean"].isin(placas_venda)].copy() if not df_manut.empty else pd.DataFrame()

    # Consolidar combustível por placa
    resumo_comb = df_v.groupby("Placa_Clean").agg(
        Modelo=("Modelo", "first"),
        Qtd_Abast=("Litragem", "count"),
        Litros=("Litragem", "sum"),
        Arla=("Litragem", lambda x: 0),  # placeholder
        Valor_Comb=("Valor total", "sum"),
    ).reset_index()

    # Tentar pegar Arla separado (tipo de combustível)
    if "Combustivel" in df_v.columns:
        arla_por_placa = df_v[df_v["Combustivel"].str.upper().str.contains("ARLA", na=False)].groupby("Placa_Clean")["Valor total"].sum()
        resumo_comb["Arla"] = resumo_comb["Placa_Clean"].map(arla_por_placa).fillna(0)

    # Consolidar manutenção por placa
    if not df_m.empty and "ValorTotal" in df_m.columns:
        resumo_manut = df_m.groupby("Placa_Clean")["ValorTotal"].sum().reset_index()
        resumo_manut.columns = ["Placa_Clean", "Valor_Manut"]
        resumo = resumo_comb.merge(resumo_manut, on="Placa_Clean", how="left")
    else:
        resumo = resumo_comb.copy()
        resumo["Valor_Manut"] = 0

    resumo["Valor_Manut"] = resumo["Valor_Manut"].fillna(0)
    resumo["Total_Geral"] = resumo["Valor_Comb"] + resumo["Valor_Manut"]
    resumo = resumo.sort_values("Total_Geral", ascending=False)

    # Cabeçalho descritivo
    ws["A1"] = "CUSTOS INFORMATIVOS — VEÍCULOS PARA VENDA"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = f"Período: {config.MES}/{config.ANO}  |  Veículos fora do escopo FKM operacional — controle de outro setor"
    ws["A2"].font = Font(italic=True)
    ws.append([])

    # Cabeçalhos da tabela
    cabecalhos = ["Placa", "Modelo", "Qtd Abast.", "Litros", "Valor Combust.", "Arla (R$)", "Manutenção (R$)", "Total Geral (R$)"]
    ws.append(cabecalhos)
    linha_cab = ws.max_row
    _estilo_cabecalho(ws, linha_cab, "C0392B")

    # Dados
    for _, row in resumo.iterrows():
        ws.append([
            row["Placa_Clean"],
            row.get("Modelo", ""),
            int(row["Qtd_Abast"]),
            round(row["Litros"], 1),
            round(row["Valor_Comb"], 2),
            round(row["Arla"], 2),
            round(row["Valor_Manut"], 2),
            round(row["Total_Geral"], 2),
        ])

    # Linha de totais
    n = len(resumo)
    linha_total = ws.max_row + 1
    ws.cell(linha_total, 1).value = f"TOTAL ({n} veículos)"
    ws.cell(linha_total, 3).value = int(resumo["Qtd_Abast"].sum())
    ws.cell(linha_total, 4).value = round(resumo["Litros"].sum(), 1)
    ws.cell(linha_total, 5).value = round(resumo["Valor_Comb"].sum(), 2)
    ws.cell(linha_total, 6).value = round(resumo["Arla"].sum(), 2)
    ws.cell(linha_total, 7).value = round(resumo["Valor_Manut"].sum(), 2)
    ws.cell(linha_total, 8).value = round(resumo["Total_Geral"].sum(), 2)
    _estilo_total(ws, linha_total)

    _ajustar_colunas(ws, {"A": 12, "B": 30, "C": 12, "D": 12, "E": 18, "F": 12, "G": 18, "H": 18})

    # Formatar colunas monetárias
    for row in ws.iter_rows(min_row=linha_cab + 1, max_row=ws.max_row):
        for col in [5, 6, 7, 8]:
            row[col - 1].number_format = 'R$ #,##0.00'
        row[3].number_format = '#,##0.0'

    return len(resumo)


def _gerar_aba_referencia(wb, df_comb, df_manut):
    """Aba com custos de veículos de unidades REFERÊNCIA (outra empresa do grupo)."""
    ws = wb.create_sheet("Referência")

    mask_comb = df_comb["Filial_Final"].apply(_eh_referencia)
    df_c = df_comb[mask_comb].copy()

    mask_manut = df_manut["FILIAL"].apply(_eh_referencia) if not df_manut.empty and "FILIAL" in df_manut.columns else pd.Series([], dtype=bool)
    df_m = df_manut[mask_manut].copy() if not df_manut.empty else pd.DataFrame()

    # Cabeçalho descritivo
    ws["A1"] = "CUSTOS INFORMATIVOS — UNIDADES REFERÊNCIA"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = f"Período: {config.MES}/{config.ANO}  |  Veículos de outras empresas do grupo — controle externo ao FKM"
    ws["A2"].font = Font(italic=True)
    ws.append([])

    # — Seção Combustível —
    if not df_c.empty:
        ws.append(["COMBUSTÍVEL"])
        ws.cell(ws.max_row, 1).font = Font(bold=True, size=11)

        cabecalhos = ["Placa", "Unidade", "Qtd Abast.", "Litros", "Valor (R$)"]
        ws.append(cabecalhos)
        _estilo_cabecalho(ws, ws.max_row, "2471A3")

        resumo_c = df_c.groupby(["Placa_Clean", "Filial_Final"]).agg(
            Qtd=("Litragem", "count"),
            Litros=("Litragem", "sum"),
            Valor=("Valor total", "sum"),
        ).reset_index().sort_values("Valor", ascending=False)

        for _, r in resumo_c.iterrows():
            ws.append([r["Placa_Clean"], r["Filial_Final"], int(r["Qtd"]), round(r["Litros"], 1), round(r["Valor"], 2)])

        linha_total_c = ws.max_row + 1
        ws.cell(linha_total_c, 1).value = f"TOTAL ({resumo_c['Placa_Clean'].nunique()} veículos)"
        ws.cell(linha_total_c, 3).value = int(resumo_c["Qtd"].sum())
        ws.cell(linha_total_c, 4).value = round(resumo_c["Litros"].sum(), 1)
        ws.cell(linha_total_c, 5).value = round(resumo_c["Valor"].sum(), 2)
        _estilo_total(ws, linha_total_c)

        for row in ws.iter_rows(min_row=ws.max_row - len(resumo_c), max_row=ws.max_row):
            row[4].number_format = 'R$ #,##0.00'
            row[3].number_format = '#,##0.0'
    else:
        ws.append(["Sem registros de combustível para unidades REFERÊNCIA no período."])

    ws.append([])

    # — Seção Manutenção —
    if not df_m.empty and "ValorTotal" in df_m.columns:
        ws.append(["MANUTENÇÃO"])
        ws.cell(ws.max_row, 1).font = Font(bold=True, size=11)

        cabecalhos = ["Placa", "Unidade", "Natureza", "Valor (R$)"]
        ws.append(cabecalhos)
        _estilo_cabecalho(ws, ws.max_row, "2471A3")

        col_natureza = "Natureza_Correta" if "Natureza_Correta" in df_m.columns else "TipoItem"
        resumo_m = df_m.groupby(["Placa_Clean", "FILIAL", col_natureza])["ValorTotal"].sum().reset_index()
        resumo_m.columns = ["Placa", "Unidade", "Natureza", "Valor"]
        resumo_m = resumo_m.sort_values("Valor", ascending=False)

        for _, r in resumo_m.iterrows():
            ws.append([r["Placa"], r["Unidade"], r["Natureza"], round(r["Valor"], 2)])

        linha_total_m = ws.max_row + 1
        ws.cell(linha_total_m, 1).value = f"TOTAL ({resumo_m['Placa'].nunique()} veículos)"
        ws.cell(linha_total_m, 4).value = round(resumo_m["Valor"].sum(), 2)
        _estilo_total(ws, linha_total_m)

        for row in ws.iter_rows(min_row=ws.max_row - len(resumo_m), max_row=ws.max_row):
            row[3].number_format = 'R$ #,##0.00'
    else:
        ws.append(["Sem registros de manutenção para unidades REFERÊNCIA no período."])

    _ajustar_colunas(ws, {"A": 12, "B": 35, "C": 25, "D": 18, "E": 18})

    return df_c["Placa_Clean"].nunique() if not df_c.empty else 0


def main():
    print("=" * 70)
    print(f"RELATÓRIO INFORMATIVO — CUSTOS FORA DO FKM — {config.MES}/{config.ANO}")
    print("=" * 70)

    # Carregar arquivos gerais do período
    pasta = os.path.join(config.DIRETORIO_BASE_SAIDA, config.PASTA_PERIODO)
    mes_num = config.obter_numero_mes()
    ano_short = config.ANO[-2:]

    arq_comb = os.path.join(pasta, f"{mes_num}{ano_short} COMBUSTIVEL GRITSCH TRANSPORTES GERAL.xlsx")
    arq_manut = os.path.join(pasta, f"{mes_num}{ano_short} MANUTENÇÃO GRITSCH TRANSPORTES GERAL.xlsx")

    if not os.path.exists(arq_comb):
        print(f"❌ Arquivo de combustível geral não encontrado: {arq_comb}")
        print("   Execute executar_resumos.py primeiro.")
        return

    print(f"\n📂 Carregando dados...")
    df_comb = pd.read_excel(arq_comb)
    df_comb["Placa_Clean"] = df_comb["Placa_Clean"].astype(str).str.upper().str.strip()

    df_manut = pd.DataFrame()
    if os.path.exists(arq_manut):
        df_manut = pd.read_excel(arq_manut)
        if "Placa_Clean" not in df_manut.columns and "Placa" in df_manut.columns:
            df_manut["Placa_Clean"] = df_manut["Placa"].astype(str).str.replace("-", "").str.strip().str.upper()
    else:
        print(f"⚠️  Arquivo de manutenção não encontrado — relatório sem dados de manutenção.")

    # Criar workbook
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # remove aba padrão vazia

    print("\n📊 Gerando abas...")

    n_venda = _gerar_aba_veiculos_venda(wb, df_comb, df_manut)
    print(f"   ✅ Veículos para Venda: {n_venda} veículos")

    n_ref = _gerar_aba_referencia(wb, df_comb, df_manut)
    print(f"   ✅ Referência: {n_ref} veículos")

    # Salvar
    nome_arquivo = f"{mes_num}{ano_short} Custos Informativos — Fora do FKM.xlsx"
    caminho_saida = os.path.join(pasta, nome_arquivo)

    wb.save(caminho_saida)
    print(f"\n✅ Relatório salvo em: {caminho_saida}")
    print("=" * 70)


if __name__ == "__main__":
    main()
