"""
Gerador de Relatório de Validação de FKMs em Excel
Cria um relatório detalhado para facilitar correção manual dos FKMs.
"""

import os
import glob
import pandas as pd
import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from src import config

# Mapeamento do nome da filial para a pasta oficial do sistema
MAPA_FILIAIS_NOME_PASTA = {
    "Gritsch Blumenau": "GRITSCH - BLN",
    "Gritsch Brasília": "GRITSCH - BSB",
    "Gritsch Campo Grande": "GRITSCH - CGR",
    "Gritsch Cascavel": "GRITSCH - CSC",
    "Gritsch Caxias do Sul": "GRITSCH - CXJ",
    "Gritsch Chapecó": "GRITSCH - CHA",
    "Gritsch Criciuma": "GRITSCH - CRI",
    "Gritsch Criciúma": "GRITSCH - CRI",
    "Gritsch Cuiabá": "GRITSCH - CGB",
    "Gritsch Curitiba": "GRITSCH - CWB BASE",
    "Gritsch Curitiba Base": "GRITSCH - CWB BASE",
    "Gritsch Curitiba Dir": "GRITSCH - CWB DIR",
    "Gritsch Curitibanos": "GRITSCH - CTB",
    "Gritsch Florianópolis": "GRITSCH - FLN",
    "Gritsch Goiânia": "GRITSCH - GOI",
    "Gritsch Guarapuava": "GRITSCH - GPA",
    "Gritsch Itumbiara": "GRITSCH - ITR",
    "Gritsch Joinville": "GRITSCH - JOI",
    "Gritsch Londrina": "GRITSCH - LDB",
    "Gritsch Matriz": "GRITSCH - MATRIZ",
    "Gritsch Maringá": "GRITSCH - MGA",
    "Gritsch Pato Branco": "GRITSCH - PBC",
    "Gritsch Petrolina": "GRITSCH - PET",
    "Gritsch PET": "GRITSCH - PET",
    "Gritsch Ponta Grossa": "GRITSCH - PGR",
    "Gritsch Palmas": "GRITSCH - PMW",
    "Gritsch Porto Alegre": "GRITSCH - POA",
    "Gritsch Rondonópolis": "GRITSCH - RDN",
    "Gritsch Rio Verde": "GRITSCH - RVD",
    "Gritsch São Paulo": "GRITSCH - SAO PERUS",
    "Gritsch São Paulo Perus": "GRITSCH - SAO PERUS",
    "Gritsch Sinop": "GRITSCH - SNO",
    "Gritsch Salvador": "GRITSCH - SSA",
}


def validar_arquivo_para_relatorio(caminho_fkm, mes_ano_fkm):
    """Valida um FKM e retorna dados estruturados para o relatório."""

    # Estrutura de retorno
    resultado = {
        "filial": None,
        "arquivo": os.path.basename(caminho_fkm),
        "status": "APROVADO",
        "placas_faltantes": [],
        "placas_extras": [],
        "erros_combustivel": [],
        "erros_manutencao": [],
        "erros_hodometro": []
    }

    try:
        xl = pd.ExcelFile(caminho_fkm)
    except Exception as e:
        resultado["status"] = "ERRO"
        return resultado

    # Encontrar aba FKM
    sheet_name = None
    for s in xl.sheet_names:
        if "fkm" in s.lower() and mes_ano_fkm in s:
            sheet_name = s
            break

    if not sheet_name:
        for s in xl.sheet_names:
            if "fkm" in s.lower():
                sheet_name = s
                break

    if not sheet_name:
        resultado["status"] = "ERRO"
        return resultado

    df_raw = xl.parse(sheet_name, header=None)

    # Identificar filial
    branch_name_raw = None
    if df_raw.shape[0] > 1 and df_raw.shape[1] > 5:
        branch_name_raw = df_raw.iloc[1, 5]

    if not branch_name_raw or pd.isna(branch_name_raw):
        base_name = os.path.basename(caminho_fkm).upper()
        for k, v in MAPA_FILIAIS_NOME_PASTA.items():
            sigla = v.split("-")[-1].strip()
            if sigla in base_name:
                branch_name_raw = k
                break

    if not branch_name_raw:
        resultado["status"] = "ERRO"
        return resultado

    branch_name = str(branch_name_raw).strip()
    resultado["filial"] = branch_name

    folder_name = MAPA_FILIAIS_NOME_PASTA.get(branch_name)
    if not folder_name:
        resultado["status"] = "ERRO"
        return resultado

    branch_dir = os.path.join(config.DIRETORIO_BASE_SAIDA, config.PASTA_PERIODO, folder_name)
    if not os.path.exists(branch_dir):
        resultado["status"] = "ERRO"
        return resultado

    # Carregar dados oficiais
    filial_acronym = folder_name.replace(" - ", "  ")
    caminho_frota = os.path.join(branch_dir, f"Frota - {filial_acronym}.xlsx")
    caminho_comb = os.path.join(branch_dir, f"Combustivel - {filial_acronym}.xlsx")
    caminho_manut = os.path.join(branch_dir, f"Manutencao - {filial_acronym}.xlsx")

    df_frota = pd.read_excel(caminho_frota) if os.path.exists(caminho_frota) else pd.DataFrame()
    df_comb = pd.read_excel(caminho_comb) if os.path.exists(caminho_comb) else pd.DataFrame()
    df_manut = pd.read_excel(caminho_manut) if os.path.exists(caminho_manut) else pd.DataFrame()

    if df_frota.empty:
        resultado["status"] = "ERRO"
        return resultado

    df_frota["Placa_Clean"] = df_frota["Placa"].astype(str).str.replace("-", "").str.strip().str.upper()

    if not df_comb.empty:
        df_comb["Placa_Clean"] = df_comb["Placa"].astype(str).str.replace("-", "").str.strip().str.upper()
    if not df_manut.empty:
        df_manut["Placa_Clean"] = df_manut["Placa"].astype(str).str.replace("-", "").str.strip().str.upper()

    # Encontrar cabeçalho do FKM
    header_row_idx = 3
    for idx in range(min(15, df_raw.shape[0])):
        row_vals = [str(x).lower().strip() for x in df_raw.iloc[idx].tolist() if pd.notna(x)]
        if "placa" in row_vals and "modelo" in row_vals:
            header_row_idx = idx
            break

    headers = df_raw.iloc[header_row_idx].tolist()
    df_fkm = df_raw.iloc[header_row_idx + 1:].copy()
    df_fkm.columns = [str(h).strip() if pd.notna(h) else f"col_{idx}" for idx, h in enumerate(headers)]

    df_fkm["Placa_Clean"] = df_fkm["Placa"].astype(str).str.replace("-", "").str.strip().str.upper()
    df_fkm = df_fkm[df_fkm["Placa_Clean"].str.len() == 7]

    # Converter colunas
    colunas_valores = ["Km Inicial", "Km Final", "Total de Km", "Litros Comb.", "Valor Comb.",
                       "Arla", "Lataria e Pintura", "Manutenção em Geral", "Rodas / Pneus"]

    for col in colunas_valores:
        if col in df_fkm.columns:
            df_fkm[col] = pd.to_numeric(df_fkm[col], errors="coerce").fillna(0)

    # Validar placas
    fkm_plates = set(df_fkm["Placa_Clean"])
    frota_plates = set(df_frota[df_frota["Placa_Clean"] != "TOTAL DE VEÍCULOS"]["Placa_Clean"])

    missing_in_fkm = frota_plates - fkm_plates
    extra_in_fkm = fkm_plates - frota_plates

    # Placas faltantes
    if missing_in_fkm:
        # Calcular custos oficiais
        custos_comb = {}
        if not df_comb.empty and "Combustivel" in df_comb.columns:
            df_comb_arla = df_comb[df_comb["Combustivel"].str.contains("Arla", case=False, na=False)].groupby("Placa_Clean").agg(
                Arla_Valor=("Valor total", "sum")
            ).reset_index()

            df_comb_no_arla = df_comb[~df_comb["Combustivel"].str.contains("Arla", case=False, na=False)].groupby("Placa_Clean").agg(
                Litros=("Litragem", "sum"),
                Valor_Comb=("Valor total", "sum")
            ).reset_index()

            for _, r in df_comb_no_arla.iterrows():
                custos_comb[r["Placa_Clean"]] = {"litros": r["Litros"], "valor": r["Valor_Comb"], "arla": 0.0}

            for _, r in df_comb_arla.iterrows():
                if r["Placa_Clean"] not in custos_comb:
                    custos_comb[r["Placa_Clean"]] = {"litros": 0.0, "valor": 0.0, "arla": 0.0}
                custos_comb[r["Placa_Clean"]]["arla"] = r["Arla_Valor"]

        custos_manut = {}
        if not df_manut.empty:
            df_manut_agg = df_manut.groupby("Placa_Clean").agg(Total_Manut=("ValorTotal", "sum")).reset_index()
            for _, r in df_manut_agg.iterrows():
                custos_manut[r["Placa_Clean"]] = r["Total_Manut"]

        for placa in sorted(missing_in_fkm):
            veiculo = df_frota[df_frota["Placa_Clean"] == placa]
            c_info = custos_comb.get(placa, {"litros": 0.0, "valor": 0.0, "arla": 0.0})
            m_val = custos_manut.get(placa, 0.0)

            resultado["placas_faltantes"].append({
                "Placa": veiculo.iloc[0]["Placa"] if not veiculo.empty else placa,
                "Modelo": veiculo.iloc[0].get("Modelo", "") if not veiculo.empty else "",
                "Litros_Oficial": c_info["litros"],
                "Valor_Comb_Oficial": c_info["valor"],
                "Arla_Oficial": c_info["arla"],
                "Manutencao_Oficial": m_val,
                "Total_Custos": c_info["valor"] + c_info["arla"] + m_val
            })
            resultado["status"] = "REJEITADO"

    # Placas extras
    for placa in sorted(extra_in_fkm):
        resultado["placas_extras"].append({
            "Placa": placa,
            "Acao": "Remover ou justificar"
        })
        resultado["status"] = "REJEITADO"

    # Validar combustível
    if not df_comb.empty and "Combustivel" in df_comb.columns:
        df_comb_arla = df_comb[df_comb["Combustivel"].str.contains("Arla", case=False, na=False)].groupby("Placa_Clean").agg(
            Arla_Oficial=("Valor total", "sum")
        ).reset_index()

        df_comb_no_arla = df_comb[~df_comb["Combustivel"].str.contains("Arla", case=False, na=False)].groupby("Placa_Clean").agg(
            Litros_Oficial=("Litragem", "sum"),
            Valor_Oficial=("Valor total", "sum")
        ).reset_index()

        df_fkm_grouped = df_fkm.groupby("Placa_Clean").agg(
            Placa=("Placa", "first"),
            Litros_FKM=("Litros Comb.", "sum"),
            Valor_FKM=("Valor Comb.", "sum"),
            Arla_FKM=("Arla", "sum")
        ).reset_index()

        df_comp = df_fkm_grouped.merge(df_comb_arla, on="Placa_Clean", how="outer")
        df_comp = df_comp.merge(df_comb_no_arla, on="Placa_Clean", how="outer").fillna(0)

        for _, r in df_comp.iterrows():
            if r["Placa_Clean"] not in fkm_plates:
                continue

            if abs(r["Litros_FKM"] - r["Litros_Oficial"]) > 0.5:
                resultado["erros_combustivel"].append({
                    "Placa": r["Placa_Clean"],
                    "Tipo": "Litros",
                    "Valor_FKM": r["Litros_FKM"],
                    "Valor_Oficial": r["Litros_Oficial"],
                    "Diferenca": r["Litros_FKM"] - r["Litros_Oficial"]
                })
                resultado["status"] = "REJEITADO"

            if abs(r["Valor_FKM"] - r["Valor_Oficial"]) > 1.0:
                resultado["erros_combustivel"].append({
                    "Placa": r["Placa_Clean"],
                    "Tipo": "Valor Combustível",
                    "Valor_FKM": r["Valor_FKM"],
                    "Valor_Oficial": r["Valor_Oficial"],
                    "Diferenca": r["Valor_FKM"] - r["Valor_Oficial"]
                })
                resultado["status"] = "REJEITADO"

            if abs(r["Arla_FKM"] - r["Arla_Oficial"]) > 0.5:
                resultado["erros_combustivel"].append({
                    "Placa": r["Placa_Clean"],
                    "Tipo": "Arla",
                    "Valor_FKM": r["Arla_FKM"],
                    "Valor_Oficial": r["Arla_Oficial"],
                    "Diferenca": r["Arla_FKM"] - r["Arla_Oficial"]
                })
                resultado["status"] = "REJEITADO"

    # Validar manutenção
    if not df_manut.empty:
        df_manut_agg = df_manut.groupby("Placa_Clean").agg(Manut_Oficial=("ValorTotal", "sum")).reset_index()

        df_fkm["Manut_Total_FKM"] = df_fkm["Lataria e Pintura"] + df_fkm["Manutenção em Geral"] + df_fkm["Rodas / Pneus"]
        df_fkm_manut = df_fkm.groupby("Placa_Clean").agg(
            Placa=("Placa", "first"),
            Manut_FKM=("Manut_Total_FKM", "sum")
        ).reset_index()

        df_comp_manut = df_fkm_manut.merge(df_manut_agg, on="Placa_Clean", how="outer").fillna(0)

        for _, r in df_comp_manut.iterrows():
            if r["Placa_Clean"] not in fkm_plates:
                continue

            if abs(r["Manut_FKM"] - r["Manut_Oficial"]) > 1.0:
                resultado["erros_manutencao"].append({
                    "Placa": r["Placa_Clean"],
                    "Valor_FKM": r["Manut_FKM"],
                    "Valor_Oficial": r["Manut_Oficial"],
                    "Diferenca": r["Manut_FKM"] - r["Manut_Oficial"]
                })
                resultado["status"] = "REJEITADO"

    return resultado


def gerar_relatorio_excel(resultados, caminho_saida):
    """Gera relatório Excel com todas as validações."""

    with pd.ExcelWriter(caminho_saida, engine='openpyxl') as writer:
        # ABA 1: Resumo
        resumo_data = []
        for r in resultados:
            resumo_data.append({
                "Filial": r["filial"],
                "Arquivo": r["arquivo"],
                "Status": r["status"],
                "Placas_Faltantes": len(r["placas_faltantes"]),
                "Placas_Extras": len(r["placas_extras"]),
                "Erros_Combustivel": len(r["erros_combustivel"]),
                "Erros_Manutencao": len(r["erros_manutencao"]),
                "Erros_Hodometro": len(r["erros_hodometro"])
            })

        df_resumo = pd.DataFrame(resumo_data)
        df_resumo.to_excel(writer, sheet_name="Resumo", index=False)

        # ABA 2: Placas Faltantes
        placas_falt_data = []
        for r in resultados:
            for p in r["placas_faltantes"]:
                placas_falt_data.append({
                    "Filial": r["filial"],
                    "Arquivo": r["arquivo"],
                    **p
                })

        if placas_falt_data:
            df_placas_falt = pd.DataFrame(placas_falt_data)
            df_placas_falt.to_excel(writer, sheet_name="Placas Faltantes", index=False)

        # ABA 3: Placas Extras
        placas_extra_data = []
        for r in resultados:
            for p in r["placas_extras"]:
                placas_extra_data.append({
                    "Filial": r["filial"],
                    "Arquivo": r["arquivo"],
                    **p
                })

        if placas_extra_data:
            df_placas_extra = pd.DataFrame(placas_extra_data)
            df_placas_extra.to_excel(writer, sheet_name="Placas Extras", index=False)

        # ABA 4: Erros Combustível
        erros_comb_data = []
        for r in resultados:
            for e in r["erros_combustivel"]:
                erros_comb_data.append({
                    "Filial": r["filial"],
                    "Arquivo": r["arquivo"],
                    **e
                })

        if erros_comb_data:
            df_erros_comb = pd.DataFrame(erros_comb_data)
            df_erros_comb.to_excel(writer, sheet_name="Erros Combustivel", index=False)

        # ABA 5: Erros Manutenção
        erros_manut_data = []
        for r in resultados:
            for e in r["erros_manutencao"]:
                erros_manut_data.append({
                    "Filial": r["filial"],
                    "Arquivo": r["arquivo"],
                    **e
                })

        if erros_manut_data:
            df_erros_manut = pd.DataFrame(erros_manut_data)
            df_erros_manut.to_excel(writer, sheet_name="Erros Manutencao", index=False)

    # Aplicar formatação
    wb = load_workbook(caminho_saida)

    # Formatar aba Resumo
    if "Resumo" in wb.sheetnames:
        ws = wb["Resumo"]

        # Cabeçalho
        for cell in ws[1]:
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            cell.font = Font(bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center")

        # Status com cores
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            status_cell = row[2]  # Coluna Status
            if status_cell.value == "APROVADO":
                status_cell.fill = PatternFill(start_color="00B050", end_color="00B050", fill_type="solid")
                status_cell.font = Font(bold=True, color="FFFFFF")
            elif status_cell.value == "REJEITADO":
                status_cell.fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
                status_cell.font = Font(bold=True, color="FFFFFF")

        # Ajustar larguras
        ws.column_dimensions["A"].width = 25
        ws.column_dimensions["B"].width = 35
        ws.column_dimensions["C"].width = 12

    wb.save(caminho_saida)


def main():
    print("=" * 80)
    print(f"📊 GERADOR DE RELATÓRIO DE VALIDAÇÃO - {config.MES}/{config.ANO}")
    print("=" * 80)

    pasta_retorno = os.path.join(config.PROJECT_ROOT, "dados", "retornados")

    if not os.path.exists(pasta_retorno):
        print(f"❌ Diretório não encontrado: {pasta_retorno}")
        return

    arquivos = glob.glob(os.path.join(pasta_retorno, "*.xls*"))
    arquivos = [a for a in arquivos if not a.endswith(".bak") and not os.path.basename(a).startswith("~$")]

    if not arquivos:
        print(f"⚠️ Nenhum arquivo FKM encontrado em dados/retornados/")
        return

    print(f"📂 Analisando {len(arquivos)} arquivos...")

    mes_ano_cod = f"{config.obter_numero_mes()}{config.ANO[-2:]}"

    resultados = []
    for arq in sorted(arquivos):
        print(f"  📄 {os.path.basename(arq)}...")
        resultado = validar_arquivo_para_relatorio(arq, mes_ano_cod)
        resultados.append(resultado)

    # Gerar relatório
    caminho_saida = os.path.join(pasta_retorno, f"Relatorio_Validacao_{config.MES}_{config.ANO}.xlsx")
    gerar_relatorio_excel(resultados, caminho_saida)

    print("\n" + "=" * 80)
    print(f"✅ RELATÓRIO GERADO COM SUCESSO!")
    print(f"📁 Arquivo: {caminho_saida}")
    print("=" * 80)
    print("\n📋 Como usar:")
    print("  1. Abra o relatório no Excel")
    print("  2. Veja a aba 'Resumo' para status geral")
    print("  3. Abra as abas de erros para ver detalhes")
    print("  4. Copie os valores oficiais e cole nos FKMs")
    print("  5. Rode novamente até todos ficarem APROVADOS")


if __name__ == "__main__":
    main()
